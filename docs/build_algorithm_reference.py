"""Build the Chinese algorithm reference using the Codex bundled Python runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'AeroPlan航线生成算法与参数说明.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
sec.left_margin = sec.right_margin = Cm(1.9)
sec.footer_distance = Cm(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.16
doc.styles['Title'].font.size = Pt(23)
doc.styles['Subtitle'].font.italic = False
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
doc.styles['Caption'].font.size = Pt(9)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('AeroPlan 算法说明  ·  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for run in footer.runs: run.font.size = Pt(8)
doc.core_properties.title = 'AeroPlan航线生成算法与参数说明'
doc.core_properties.subject = '输入输出定义 计算流程 复算示例 实现边界'
doc.core_properties.author = 'AeroPlan 项目'

def p(text, style=None):
    return doc.add_paragraph(text, style)
def h(text): doc.add_heading(text, 2)
def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)
def eq(text):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(7)
    math = OxmlElement('m:oMath')
    run = OxmlElement('m:r'); txt = OxmlElement('m:t'); txt.text = text
    run.append(txt); math.append(run); para._p.append(math)
def code(text):
    para = p(text)
    para.paragraph_format.line_spacing = 1.05
    for run in para.runs: run.font.name = 'Consolas'; run.font.size = Pt(9)
def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el = OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    props.append(borders)
    for c,w in zip(t.columns,widths): c.width = Cm(w)
    for i, values in enumerate([headers] + rows):
        row = t.rows[0] if i == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        keep = OxmlElement('w:cantSplit'); trpr.append(keep)
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width = Cm(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr(); margins = OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),'75'); el.set(qn('w:type'),'dxa'); margins.append(el)
            tcpr.append(margins)
            fill = OxmlElement('w:shd'); fill.set(qn('w:fill'),'DBE5EE' if i == 0 else ('F4F7FA' if i%2 == 0 else 'FFFFFF')); tcpr.append(fill)
            para=cell.paragraphs[0]; para.paragraph_format.space_after=Pt(1); para.paragraph_format.line_spacing=1.08
            run=para.add_run(str(value)); run.font.size=Pt(9)
            run.bold=(i==0)
    p('').paragraph_format.space_after=Pt(0)
    return t

p('AeroPlan 航线生成算法\n与参数说明', 'Title')
p('城市正射影像采集预规划技术文档', 'Subtitle')
p('技术版本 terrainAGL v2  |  更新日期 2026年9月20日')
p('本项目将用户绘制的单个多边形作业区转换为有序拍摄航段、拍照位置和区内连接路径。算法采用局部米制投影、相机视场角估算、扫描线裁剪、平面覆盖补线及可视图最短路径，支持凹多边形和内部排除区。')
p('本文用于理解、复算和维护当前项目。v2 增加 4D 正射三向智能摆拍的采集意图，并将航高改为固定离地高度 AGL。水平路线生成后，对全部航段和拍照站点采样在线地形，得到逐点绝对高程。')
h('适用范围')
p('当前结果仍是采集预规划。足迹和重叠率按水平地面近似；在线 DEM 提供地形高程，不包含完整建筑 DSM。三向模式按用户角度建立近似姿态，不虚构原生曝光时序，也不以未经证实的侧向覆盖减少航线。结果不等于可直接下发机场的任务。')
table(['阅读内容','对应章节'],[
 ['接口与参数','第 2 至 3 章  输入定义 相机预设 坐标与高度'],
 ['核心计算','第 4 至 9 章  几何校验 间距 扫描 航向 补线 连接'],
 ['结果解释','第 10 至 13 章  三向姿态 地形高程 统计与导出'],
 ['复算与维护','第 14 至 16 章  算例 错误边界 代码与测试']], [4.2,13])
h('完整处理链')
p('参数与多边形校验 → 局部投影 → 下视足迹和间距 → 航向选择与扫描 → 平面补线 → 区内连接 → 采集站点与模式意图 → 路线加密 → 在线地形采样 → 逐点赋高 → 三维距离和时间统计。')
p('高度只由“地形高程＋设定离地高度”确定。机场位于楼顶不会改变区内间距或逐点高程；机场进出、返航和楼顶障碍物仍须单独规划。')

page('2 输入接口与规划参数')
code('planMission(ringLL, options = {}, holesLL = [])')
p('ringLL 是外边界经纬度数组；holesLL 是内部排除区数组；options 是参数对象。坐标顺序始终为 [经度, 纬度]，单位为度。坐标点可附带第三个分量，但算法清洗时只保留前两个分量。首尾可闭合，也可不重复首点。')
table(['字段','默认值','单位与引擎约束'],[
 ['camera','4D','字符串  4D / 4TD12 / 4TD48'],
 ['captureMode','nadir','nadir 下视 / smartOrtho 三向意图  后者仅限 4D'],
 ['altitude','100','m  固定离地高度 AGL  0 < H ≤ 1000'],
 ['smartPitch','−62.5','°  左右侧视俯仰  −65 至 −60'],
 ['smartYaw','27.5','°  左右相对偏航幅度  25 至 30  横滚固定 0'],
 ['terrainSampleSpacing','30','m  航线最大水平采样间距  5 至 100'],
 ['frontOverlap','85','%  航向重叠率  50 至 95'],
 ['sideOverlap','80','%  旁向重叠率  50 至 95'],
 ['speed','6','m/s  0 < speed ≤ 25'],
 ['heading','0','°  任意有限数  使用前归一化到 [0,180)'],
 ['autoHeading','true','布尔值  是否搜索 12 个候选方向'],
 ['crossGrid','false','布尔值  是否增加正交的第二遍扫描'],
 ['turnSeconds','3','s  每两个相邻拍摄航段间的附加时间  0 至 60']], [4.1,2.2,10.9])
p('数值必须是有限的 JavaScript Number，不能用数值字符串、NaN 或 Infinity。两个开关必须是真正的布尔值。传入部分 options 时，其余字段沿用默认值；未知扩展字段目前不会参与核心计算。')
h('兼容与两阶段计算')
p('buildingHeight、dockHeight、clearance 为旧版本参数，读取旧方案时不再参与计算；屋顶 GSD、建筑净空和相对机场高度统计已移除。目标 GSD 仍由 AGL 反算。4TD 不能选择 smartOrtho。')
code('horizontal = planMission(ringLL, options, holesLL)\nplan = await attachTerrainToPlan(\n  horizontal, options, sampleHeights, metadata)')
p('planMission 返回 terrain.status="pending" 的二维中间结果。attachTerrainToPlan 接收按输入顺序返回高程数组的异步采样函数，全部高程有效才返回 ready；网页仅允许导出 ready 的完整方案。')

page('3 相机预设与坐标高度约定')
table(['相机键','照片像素','对角视场角','参考最短间隔'],[
 ['4D','5280 × 3956','84°','0.5 s'],
 ['4TD12','4032 × 3024','82°','0.7 s'],
 ['4TD48','8064 × 6048','82°','0.7 s']], [2.8,5.4,3.8,5.2])
p('以上为代码内置预设，最短间隔用于 JPEG 模式参考提示。采用标称对角视场角和像素长宽比估算成像几何，未使用标定焦距、主点、畸变参数或实际裁切尺寸。4TD 的两种预设具有相同足迹，48 MP 模式的名义 GSD 为 12 MP 模式的一半；这不等同于实际有效分辨率保证。')
h('经纬度与局部平面')
p('输入输出使用 CGCS2000 经纬度语义，底图使用 EPSG:4490。几何计算先通过 proj4 转到 GRS80 椭球上的局部等距方位投影 AEQD，单位为米。投影中心是外边界顶点经纬度的算术平均值，不是面积重心。')
code('+proj=aeqd +lat_0=中心纬度 +lon_0=中心经度\n+ellps=GRS80 +units=m +no_defs')
p('局部坐标 x 表示东西方向，y 表示南北方向。航向按局部投影北方向顺时针计量：0° 为南北向，90° 为东西向。局部投影距离用于面积、间距和路径长度，不能理解为在所有地点均无误差的大地测量距离。')
h('离地高度与绝对高程')
eq('z飞行(经度,纬度) = z地形(经度,纬度) + altitude')
p('altitude 是飞机相对下方地形的垂直距离，不是相对机场高度，也不是一个全区不变的绝对高程。z飞行随地形变化，保留地形源原始垂直基准。在线源是多源地形高程，未统一转换为椭球高，不能标注成已校准的 EGM96 或 CGCS2000 高程。')
p('EPSG:4490 约定水平经纬度，不能单独说明第三个坐标的垂直基准。地形原始高程在 Cesium 中仅作三维位置近似显示；测绘或飞行执行前需完成垂直基准转换和现场核对。')
h('界面范围与引擎范围')
p('网页控件可比引擎校验更窄；程序调用以第 2 章约束为准。采样间距是路线加密尺度，不是 DEM 分辨率或精度。减小间距只能增加采样位置，不能恢复地形源中不存在的楼宇、细线或地形细节。')

page('4 多边形清洗与合法性校验')
h('第一步 清洗坐标')
p('每个点必须含有限经纬度，经度绝对值不超过 180°，纬度绝对值小于 85°。若首尾点的经纬度欧氏差小于 10⁻¹¹，则删除末尾重复点。每个环清洗后至少需要 3 个点，外环与所有孔洞的顶点总数最多 180。')
h('第二步 控制局部投影范围')
p('外环经度跨度不能超过 5°，因此不支持跨日期变更线的单一区域。投影后，外环 x 和 y 两方向的跨度分别不得超过 30 km。这是轴向跨度限制，不是面积或对角线长度限制。')
h('第三步 检查简单多边形')
p('每条边长度至少为 0.05 m；相邻边不得回折重叠；非相邻边不得相交或接触。每个环的面积至少为 1 m²。程序采用射线奇偶规则进行点包含判断，边界点单独标识为 0，内部为 1，外部为 −1。')
h('第四步 检查排除区')
p('每个孔洞必须严格位于外环内，不能与外环相交或接触。孔洞之间不能相交、接触或相互包含。外环和孔洞都不要求固定顺时针或逆时针排列。当前接口表示一个外环加多个孔洞，不表示多个互不相连的作业区。')
h('第五步 计算净作业面积')
eq('A环 = |Σ (xᵢ yᵢ₊₁ − xᵢ₊₁ yᵢ)| / 2')
eq('areaM2 = A外环 − Σ A孔洞')
p('这是投影平面上的鞋带公式面积，单位为平方米；不是地表起伏面积。几何处理中主要使用 EPS = 10⁻⁶ m 量级的容差，并对叉积按线段长度缩放；容差服务于数值稳健性，不代表测绘精度。')
code('validatePolygon(ringLL, holesLL)\n成功  { valid: true, errors: [], areaM2 }\n失败  { valid: false, error: "原因", errors: ["原因"] }')
p('planMission 遇到非法边界或非法参数时直接抛出 Error；validatePolygon 只检查几何并返回结果对象。')

page('5 相机足迹与拍摄间距计算')
p('以下模型用于下视摄影，也作为三向意图的保守布线基线。假设相机垂直向下、地面局部水平，照片长边横跨航线、短边沿航线。设像素宽高 Nw、Nh，对角视场角 α，固定离地高度 H=altitude。')
h('步骤一 求单位高度对应的像素尺度')
eq('dpx = √(Nw² + Nh²)')
eq('k = 2 tan(απ / 360) / dpx')
p('k 的量纲为每米高度对应的米每像素。απ/360 将半视场角转换为弧度。先由对角视场角得到对角足迹，再按像素长宽比分配为宽和长。')
h('步骤二 求地形表面名义分辨率')
eq('gsdCm = 100 H k')
p('单位为 cm/px。这是下视相机在水平地面的名义值，不是地形坡面每个位置的真实 GSD，也不代表三向摆拍侧视影像的 GSD。系统不再减去统一建筑高度。')
h('步骤三 求下视足迹和目标间距')
eq('W = H k Nw     L = H k Nh')
eq('lineSpacingM = W (1 − sideOverlap / 100)')
eq('shotSpacingM = L (1 − frontOverlap / 100)')
p('W 是横跨航线的足迹宽度，L 是沿航线的足迹长度。两个 spacing 输出是用于生成航线和拍照点的目标最大间距。为均匀排布并包含端点，实际行距和点距通常更小。')
h('参数变化的影响')
p('提高重叠率会减小间距并通常增加站点数；提高 AGL 会扩大足迹、增大名义 GSD。选择三向模式保留相同行距与站距，增加三向采集意图和预算照片数；不承诺复刻 DJI 原生稀疏航线效率。速度影响时间，通常不改变站点位置。')

page('6 扫描线裁剪与往复式排序')
h('步骤一 将多边形旋转到扫描坐标')
p('设航向 θ 已转换为弧度。u 沿飞行方向，v 垂直于飞行方向。此变换为正交坐标变换，其逆变换形式相同。')
eq('u = x sinθ + y cosθ     v = x cosθ − y sinθ')
eq('x = u sinθ + v cosθ     y = u cosθ − v sinθ')
h('步骤二 均匀布置扫描行')
eq('n = max(1, ceil((vmax − vmin) / lineSpacingM))')
eq('Δv = (vmax − vmin) / n')
eq('vⱼ = vmin + (j + 0.5) Δv     j = 0 … n−1')
p('vmin 和 vmax 取外环顶点的极值。第一行距外包范围的一侧为半个实际行距，最后一行距另一侧同样为半行距。这里的 n 是扫描行数；一个扫描行穿过凹区或孔洞时，可形成多个拍摄航段。')
h('步骤三 求扫描线与所有边的交点')
eq('u交点 = uₐ + (vⱼ − vₐ)(uᵦ − uₐ) / (vᵦ − vₐ)')
p('仅当 vₐ ≤ vⱼ < vᵦ 或 vᵦ ≤ vⱼ < vₐ 时计入交点。半开区间避免扫描线穿过顶点时重复计数，与扫描线平行的边不参与上述求交。')
p('将外环和所有孔洞的交点一起按 u 排序，并按第 1–2、第 3–4 个交点依次成对。奇偶规则保留作业区内部区间，同时剔除孔洞和凹陷外部。长度不大于 EPS 的区间被忽略。')
h('步骤四 形成往复式顺序')
p('扫描行按 v 从小到大排列。第 0 行从较小 u 飞向较大 u，第 1 行反向，之后逐行交替。奇数行不仅反转每个航段的端点，也反转同行多段的访问顺序。此阶段只确定顺序，跨段连接随后单独计算。')
code('一行交点  a < b < c < d\n保留区间  [a,b] 与 [c,d]\n中间区间  (b,c) 属于孔洞或区域外部')

page('7 自动航向与双网格计算')
h('手动航向')
eq('heading = ((输入角度 mod 180) + 180) mod 180')
p('180° 周期描述航线轴向，往返航段的实际行进方向由排序决定。因此 heading 不是每个航段的飞机机头角，接口也没有输出逐点 yaw 指令。')
h('自动航向')
p('autoHeading 为 true 时，依次测试 0°、15°、30°，直到 165°，共 12 个候选方向。每个候选先生成常规扫描航段，再计算近似时间代价，选择严格小于当前最优代价的方案；完全相等时保留先遇到的方向。')
eq('C(θ) = (Σ L拍摄 + Σ D端点直连) / speed')
eq('             + max(0, 航段数 − 1) × turnSeconds')
p('D端点直连是相邻拍摄航段之间的欧氏直线距离。评分阶段未进行孔洞绕行，也未加入后续补线，因而评分与最终路线时间可能不同。此搜索是离散启发式选择，不是全局最短航线求解。')
h('双网格')
p('crossGrid 为 true 时，每个候选同时生成 θ 和 (θ + 90°) mod 180° 两组扫描航段，先完成第 1 遍，再完成第 2 遍。两遍合并后参与代价计算；最终会分别补线，并计算从第一遍末点到第二遍首点的区内连接。')
table(['情况','当前实现'],[
 ['风向与侧风','未进入候选评分'],
 ['机场位置与出入方向','未参与起点或终点选择'],
 ['补线与孔洞绕行','选定方向后处理，未重新搜索航向'],
 ['连接路径','最终路径按允许区域校核'],
 ['两遍同位置照片','独立保留，不跨遍去重']], [5,12.2])
p('双网格提供两个方向的平面采集结构，但没有计算真实地物可见性和多视交会条件。需要真正射成果时，不能仅凭开启双网格认定建筑遮挡问题已经解决。')

page('8 平面足迹覆盖检查与补线')
p('常规均匀扫描可能漏过位于两行之间的细长突出部。completeFlatCoverage 在选定方向后检查近似照片足迹的并集，发现遗漏时补充扫描行；双网格的每一遍独立执行此过程。')
h('步骤一 构造每个航段的覆盖矩形')
eq('R = [umin − L/2, umax + L/2] × [v行 − W/2, v行 + W/2]')
p('L 和 W 来自第 5 章 AGL 下视足迹。采集站点包含航段两端，且沿线间距不大于足迹长度，因此在固定水平面、固定朝向模型内，照片足迹并集是连续矩形。它可以延伸到飞行边界或孔洞之外。')
h('步骤二 建立横向扫描事件')
p('事件包括全部多边形顶点的 v 坐标，以及落在区域纵向范围内的覆盖矩形上下边。排序后，对每个宽度大于 2EPS 的相邻事件区间，检查中点、下边上方和上边下方三个位置。靠近边界的偏移量取 min(区间宽度/4, 10EPS)。')
h('步骤三 判定尚未覆盖的区间')
p('在检查位置 v 上重新求多边形内部区间。取该位置穿过的覆盖矩形，按左端点排序并求其区间并集。若任一作业区间中出现超过 EPS 的空隙，则判定该检查位置存在平面足迹遗漏。')
h('步骤四 插入补充扫描行并迭代')
p('发现遗漏后，在该 v 处生成完整扫描行，包括这一行的全部合法区间，而不是只生成遗漏的一个短段。立即将新矩形加入覆盖集合，并在下一轮重建事件。直到无补线产生，再统一重新执行往复式排序。')
p('补线行数每遍最多 360，循环索引为 0 至 180；超过航段上限、无法生成有效补线或不能在有限轮次内完成时，返回错误，要求简化或拆分区域。coverageAddedRows 记录新增扫描行数，不是新增航段数。')
h('覆盖结论的边界')
p('此过程针对带数值容差的近似平面足迹至少一次覆盖。它没有验证边缘所需的多张影像重叠、立体重建质量或建筑遮挡，也没有外扩作业边界补拍。排除区限制飞机路径，不表示相机画面不会拍到该区域。')

page('9 区内连接与孔洞绕行')
p('相邻拍摄航段之间不直接无条件连线。createNavigator 在外边界内、孔洞内部之外寻找连接折线。外边界和孔洞边界允许经过，当前没有向内或向外施加水平安全偏移。')
h('步骤一 判断直线可见性')
p('先检查两个端点均处于允许区域，再求连接线段与全部边界边的交点参数 t。把 t=0、t=1 和交点参数排序，对每个非退化子区间的中点检查区域归属。只要有中点位于外环之外或孔洞内部，这条连接线就不可直接使用。')
h('步骤二 建立边界顶点可视图')
p('若端点可直连，立即返回两个点。否则以外环和所有孔洞顶点为节点，逐对检查可见性；可见的节点对建立无向边，边权为米制欧氏距离。边界基础图按一次任务缓存，不对每条连接重复构建。')
h('步骤三 接入起终点并运行最短路径')
p('把前一拍摄航段终点和后一航段起点加入图，与可见的边界顶点相连。Dijkstra 算法每轮在线性数组中选择未访问的最低成本节点，松弛其邻边，最后通过前驱节点反向还原连接折线。找不到路径则报错。')
code('navigate(上一航段.end, 下一航段.start)\n可直连 → [start, end]\n不可直连 → 可视图 → Dijkstra → [start, 转折点…, end]')
h('最短的具体含义')
p('这里求的是固定两端点之间、二维允许区域内的折线路径。在零缓冲多边形模型下，可视图适用于绕凹角与孔洞；它不重新排列拍摄航段，也不优化整次任务的访问顺序。')
h('工程限制')
p('飞机可在结果中贴边经过并在多边形顶点急转。算法未约束最小转弯半径、加速度、侧向净空或三维障碍物。起降场进出路线与返航路线不属于 connections。大量顶点会增加两两可见性测试开销，因此项目限制总顶点数并在 Web Worker 内计算。')
p('基础图在朴素实现下约需 O(V³) 几何工作量：O(V²) 对节点，每对可见性检查扫描边界。实际耗时还取决于航段、补线和连接数量，此量级描述不是性能承诺。')

page('10 三向近似姿态与足迹')
p('captureMode="smartOrtho" 仅允许 4D。该模式按用户确认的近似角度生成相机模型，不将近似角度冒充 DJI 固件指令。中间向下，左右使用偏航角，横滚均为 0°。')
table(['方向','俯仰角','相对航段偏航','横滚'],[
 ['左侧','−62.5°  可调 −65° 至 −60°','−27.5°  幅度可调 25° 至 30°','0°'],
 ['中间','−90°','0°','0°'],
 ['右侧','−62.5°  可调 −65° 至 −60°','+27.5°  幅度可调 25° 至 30°','0°']], [2,6.2,7,2])
h('射线与参考平面求交')
p('设右向为 x，航段前向为 y，向上为 z；俯仰 p 相对水平面，偏航 q 顺时针。相机前向 f、右向 r、上向 u 定义如下。角度运算先转弧度。')
eq('f = (sin q cos p, cos q cos p, sin p)')
eq('r = (cos q, −sin q, 0)')
eq('u = (−sin q sin p, −cos q sin p, cos p)')
p('令 tx=kNw/2、ty=kNh/2。四个角点分别取 a、b 为 ±1：射线方向 d=f+a tx r+b ty u。若 dz≥0，射线到达或越过地平线，拒绝有限足迹计算；其余角点取 t=−H/dz，得到水平足迹偏移 (t dx,t dy)。')
p('captureModel.views 保存每向的 pitch、yawOffset、roll、四角偏移 corners 和 areaM2。各采集请求保存实际航段方向 flightHeadingDeg；全局观测方向为航段方向加相对偏航。返程会随航段转向，不固定朝地理东西方向。')
h('覆盖与照片预算')
p('三向模式仍按第 5 章下视足迹确定行距与站距，并按下视足迹检查平面遗漏。侧向足迹不参与扩大行距，因此与相同参数的垂直模式保持同一条路线。每站预算 3 张，总预算为 stationCount×3；实际曝光位置、先后顺序和摆动时序仍由原生任务决定。')
p('网页开启拍照点时展示首站三向足迹。足迹位于该站地形高程对应的水平面上，没有逐角点与坡面求交；它是近似示意，不能证明真实地形或建筑表面的重叠率。大疆作业指南提示正射摆拍可能降低最终正射质量。')

page('11 在线地形采样与逐点赋高')
h('开放地形来源')
p('采用 Mapzen Terrain Tiles 的 AWS Open Data Terrarium 瓦片，无需新增账号密钥。显示与数字采样使用同一数据源。固定采样层级为 13，地图地形最高也使用 13 级。水平索引为 Web Mercator；高度沿用多源正高基准，不转换为椭球高。')
code('https://s3.amazonaws.com/elevation-tiles-prod/terrarium/{z}/{x}/{y}.png')
eq('地形高程 m = R × 256 + G + B/256 − 32768')
p('每张瓦片 256×256 像素。解码前关闭图像色彩空间转换；透明像素、异常尺寸或异常高程直接拒绝。按像素中心位置进行跨瓦片双线性插值；渲染网格重采样到 257×257 个顶点，邻瓦片共享相同边界高度。13 级像素间隔不是原始 DEM 的精度保证。')
h('第一步 沿全部路线加密')
p('在原始 AEQD 米制坐标中，对每个拍摄航段和连接折线分段加密，保持全部采集站和绕行拐点。每个子段的水平长度不大于 terrainSampleSpacing，默认 30 m，可设 5–100 m。不会以直接经纬度插值替代原有许可路径。')
eq('加密间隔数 = max(1, ceil(子段水平长度 / 采样间距))')
h('第二步 去重采样并计算飞行高程')
p('经纬度按 12 位小数建立缓存键，同一位置只请求一次。最多 50000 个唯一点，每批 256 点；数据源内部缓存最多 128 张瓦片，最多 6 路网络请求，单次请求 15 秒超时并重试一次。缺失高程或网络失败立即终止，不回退为 0 m。')
eq('z飞行ᵢ = z地形ᵢ + altitude')
p('赋高后，legs.positions、legs.start/end/photos、connections.positions、photos、path 和 captureRequests.position 均使用 [经度,纬度,飞行绝对高程]。terrain.status 从 pending 变为 ready，附带 source、verticalDatum、sampleLevel、sampleSpacingM 和 sampleCount。')
h('采样与真实连续航迹')
p('固定 AGL 仅在采样点按数据源成立。点间是三维直线，可能跨过未采到的尖峰；DEM 缺少的建筑和树木不会自动出现。减小采样间距也不能恢复数据源中缺失的细节。')

page('12 采集站与三维任务统计')
h('采集站插值')
p('对每个水平拍摄航段长度 S，K=max(1,ceil(S/shotSpacingM))，按 Pj=P起+(j/K)(P终−P起) 生成 K+1 个站，j=0…K。两端都保留，即使很短的航段也至少两个站。各站不去重，区内连接不另设采集站。')
eq('实际站距 = S/K ≤ shotSpacingM')
p('stationCount=photos.length。垂直模式 photoCount=stationCount；三向模式 photoCount=3×stationCount，photoCountIsEstimate=true。photos 在三向模式表示采集站，不能把它理解为三张实际曝光的共同精确位置。')
h('距离与时间')
eq('d三维 = √(Δx² + Δy² + Δz²)')
eq('distanceM = Σ d拍摄三维 + Σ d连接三维')
eq('durationSeconds = distanceM / speed')
eq('                   + max(0, legCount−1) × turnSeconds')
p('horizontalDistanceM 另存同一条路线的水平距离。speed 表示沿三维折线的统一速度。附加时间按相邻拍摄航段计一次，不按绕行顶点次数累计。未计原生三向摆动、稳定和曝光周期，也不含机场往返、起降与电池分架次。')
h('地形坡度与垂直速度提示')
eq('坡度 = Δz / √(Δx² + Δy²)')
eq('所需垂速 = speed × Δz / d三维')
p('分别统计最大上升和下降坡度、所需爬升及下降速度、累计升降高度。坡度超过 20% 时提示复核，此阈值只是提示条件，不是机型安全界限；当前没有按飞机性能自动减速。')
table(['统计字段','单位与意义'],[
 ['terrainMinM / terrainMaxM','m  全部路线采样点的地形高程范围'],
 ['absoluteHeightMinM / absoluteHeightMaxM','m  地形高程范围加 AGL'],
 ['aglM / gsdCm','m / cm/px  离地高度与下视名义 GSD'],
 ['footprintWidthM / footprintLengthM','m  下视参考足迹宽长'],
 ['maxRequiredClimbRateMps / maxRequiredDescentRateMps','m/s  按三维航速推算的最大垂速']], [10.1,7.1])

page('13 输出数据与导出格式')
table(['顶层字段','含义'],[
 ['heading / heightMode / captureMode','主航向 / terrainAGL / nadir 或 smartOrtho'],
 ['projectionOrigin','原 AEQD 中心  用于保持加密几何一致'],
 ['legs','拍摄段  id pass start end photos positions flightHeadingDeg'],
 ['connections / path','加密并带高程的区内连接 / 完整路线'],
 ['photos','带高程的有序采集站  三向模式不等于实际曝光清单'],
 ['captureModel','近似姿态源  参考平面  三向相机足迹模板'],
 ['captureRequests','每站位置 序号 航段编号 飞行方向 模式 方向数'],
 ['terrain','ready 状态  来源  高程基准  数据层级与采样信息'],
 ['stats / warnings','距离 时间 数量 高程与坡度等统计 / 使用限制']], [4.9,12.3])
p('legs.id 从 1 编号；pass 为 1 或 2；captureRequests.stationIndex 从 0 编号。nativeAngles 和 nativeTiming 保持 null，表示没有生成 DJI 原生动作。captureModel.views 中的角度是用户近似模型，二者不能混淆。')
h('JSON')
p('外层 format="aeroplan"、version=2、crs="EPSG:4490"、planningStatus="terrain-sampled-unverified"。包括任务状态、options、plan、heightReference 和 notes。heightReference 明确 terrainAGL 公式、verticalDatum 以及 ellipsoidConversionApplied=false。机场坐标仅作位置标记。')
h('CSV')
p('文件名以“采集站高度.csv”结尾，UTF-8 BOM 编码，一站一行。列为采集站序号、经度_CGCS2000、纬度_CGCS2000、地形源高程_m、飞行绝对高程_m、离地高度_m、采集方式、每站预算照片数和高程基准。经纬度保留 9 位小数，高程保留 3 位。')
p('导出前要求 terrain.status="ready"、高程基准非空、所有采集站高度有效。正在重算或地形失败时，旧方案不会继续供导出。JSON 可重新导入并重算；这些文件仍不是 WPML 或 KMZ 可执行飞行任务。')

page('14 可复算的更新算例')
p('在 [116.39°,39.9°] 的局部投影中构造 300×200 m 矩形，默认参数，只将 autoHeading=false、heading=90°。无孔洞，单网格。以下平面高程是可重复的测试输入，不冒充真实在线测区高程。')
code("const f = createLocalProjection([[116.39,39.9]]);\nconst ring = [[0,0],[300,0],[300,200],[0,200]].map(f.inverse);\nconst options = {autoHeading:false,heading:90};\nconst horizontal = planMission(ring,options);\nconst plan = await attachTerrainToPlan(horizontal, options,\n  async points => points.map(() => 50),\n  {source:'测试平地',verticalDatum:'source-orthometric'});")
h('计算步骤')
p('① AGL=100 m；4D 像素 5280×3956、对角 FOV=84°。k≈0.0002729489325，名义 GSD≈2.729489 cm/px；下视足迹宽≈144.117036 m，长≈107.978598 m。')
p('② 旁向 80%、航向 85%：行距≈28.823407 m，站距≈16.196790 m。横向需要 ceil(200/28.823407)=7 行，实际行距约 28.571 m。')
p('③ 每行 K=ceil(300/16.196790)=19，包含两端共 20 站，实际站距约 15.789474 m。7 行合计 140 站。垂直模式预算 140 张；三向模式预算 420 张，水平路线相同。')
p('④ 测试地形每点 50 m，则所有飞行绝对高程为 150 m。三维距离等于水平距离；在线地形起伏时，需要按逐点 Δz 重算。')
table(['结果','平地代码实算值'],[
 ['面积 / 航段 / 连接','59999.999541 m² / 7 段 / 6 条连接'],
 ['拍摄距离 / 连接距离','2099.999984 m / 171.433621 m'],
 ['区内总距离','2271.433605 m'],
 ['区内预计时间','2271.433605/6 + 6×3 = 396.572268 s\n约 6.609538 min  不含三向摆动周期']], [6.2,11])
h('斜坡的独立验算')
p('若一条 100 m 东向水平航段的地形由 0 m 线性上升至 50 m，AGL=100 m，则飞行绝对高程从 100 m 升至 150 m；三维长度 √(100²+50²)=111.803399 m。三维航速 6 m/s 时行进约 18.633900 s，所需爬升约 2.683282 m/s。')

page('15 错误处理与适用边界')
table(['检查项','拒绝或提示条件'],[
 ['几何与参数','自交 接触 过近点 孔洞非法 非有限数或参数越界'],
 ['机型与角度','4TD 不接受 smartOrtho  近似俯仰与偏航超范围则拒绝'],
 ['容量','总顶点 180  轴向跨度 30 km  航段 4000\n预算照片 15000  唯一地形点 50000'],
 ['地形失败','缺瓦片 缺高程 无效解码 超时  不以 0 替代'],
 ['过快拍照','站间时间不足  提示复核速度和原生摆拍周期'],
 ['地形过陡','按采样点提示坡度与所需垂速  不自动减速'],
 ['补线或连接失败','要求简化边界 调整排除区或拆分区域']], [4.8,12.4])
h('本版本完成的计算')
p('支持任意合法凹多边形和内部孔洞，完成水平扫描、补线、区内绕行、用户近似三向姿态、足迹示意、在线 DEM 采样与带高程路线。高程采样与地图地形共享同一开放数据源。')
h('当前不作出的保证')
p('固定 AGL 不等于跨越建筑的净空。地形源不是完整 DSM，模型不包含线缆、树木或机场屋顶避障；当前也不验证每张倾斜影像与真实坡面的交点、街巷可见性或真正射成果完整性。')
p('三向足迹与预算是用户给定姿态下的工程近似；原生任务是否接受这些姿态、怎样摆动及曝光，仍需 DJI 任务文件和现场验证。精确飞行还需要坐标与垂直基准转换、转弯和爬升能力、机场进出、风、电量和分架次。')
h('坐标与公开服务')
p('天地图 EPSG:4490 底图继续保留；地形瓦片按 WGS84 Web Mercator 索引，未进行测绘级水平基准转换。导入仍只接受单个 GeoJSON Polygon 的 4490 或 4326 经纬度，不接受直接混入 GCJ-02、BD-09 或墨卡托米制坐标。')
p('在线服务可用性与原始数据精度会变化。网页显示采样状态和来源，失败时允许重试。地形源正高在 Cesium 中作近似显示，不能当作已转换的椭球高导入飞机。')

page('16 实现定位与验证方法')
table(['源码','职责'],[
 ['src/planner.js','参数校验 局部投影 扫描 補线 最短连接与采集站'],
 ['src/capture-geometry.js','用户近似姿态  相机射线与水平面求交'],
 ['src/terrain-source.js','Terrarium解码 跨瓦片插值 缓存与Cesium地形'],
 ['src/terrain-route.js','米制路线加密 批量采样 逐点赋高 三维统计'],
 ['src/mission-export.js','JSON v2 与采集站高度 CSV  导出条件检查'],
 ['src/planner.worker.js','异步水平计算  返回任务编号和结果或错误'],
 ['src/main.js / src/map.js','机型模式交互  异步地形流程  地图足迹与预览'],
 ['docs/algorithm-reference.test.js','本文矩形与斜坡算例回归验证']], [6.5,10.7])
p('网页参数修改采用 220 ms 防抖。任务编号贯穿水平计算和地形采样；旧任务后返回时不覆盖新结果。重算立即停用旧导出，地形全部采样成功后才恢复。')
h('验证内容')
p('测试包含多边形与孔洞路径、相机模式限制、左右近似足迹对称性、AGL 与旧参数兼容、固定坡面三维距离、绕行峰值加密、采样限额与失败处理、Terrarium 小数及负高程解码、瓦片接缝、并发缓存和导出高程字段。')
code('npm test\nnpm run build\nnpm run test:deployment')
p('自动测试约束实现行为，不替代测绘精度检验、飞行性能验证或现场试飞。修改参数或高程流程后应同步重算本文示例。')
h('资料与访问入口')
p('项目  https://github.com/marky789/aeroplan-ortho-planner')
p('网页  https://marky789.github.io/aeroplan-ortho-planner/')
p('地形数据登记  https://registry.opendata.aws/terrain-tiles/')
p('编码与版权  https://github.com/tilezen/joerd/tree/master/docs')
p('相机规格  https://enterprise.dji.com/dock-3/specs')
p('三向支持与质量提示  DJI Dock 3 Matrice 4D 系列作业指南第 55 页，可从大疆机场 3 官方下载中心获取。侧向 −62.5°、偏航 ±27.5°、横滚 0° 来自本项目用户确认的近似建模参数。')

doc.save(OUT)
print(OUT)
